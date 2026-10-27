"""Goods profile: group listings by model and memory, compute market prices.

Avito titles in structured categories look like "iPhone 15, 128 ГБ, SIM +
eSIM": the first part is the model, one part is the storage. A group is
"model + storage"; its market price is the median after dropping extreme
prices, which are usually broken devices, parts, scams or typos.
"""

from __future__ import annotations

import re
import statistics
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import apartments

MEMORY_PATTERN = re.compile(r"^\d+\s*(?:ГБ|ТБ|GB|TB)$", re.IGNORECASE)
# Prices outside [LOW, HIGH] x raw median are left out of the market price.
OUTLIER_LOW = 0.5
OUTLIER_HIGH = 1.8
CHEAP_SHARE = 0.85
MIN_GROUP_SIZE = 3
# A "private" seller with this many closed listings is a shop or a reseller.
SHOP_CLOSED_LISTINGS = 100
# A reseller network posing as private people: fresh accounts that list the
# same model at one exact, non-round price again and again ("34 990"), or
# "sealed, with warranty" phones long out of production. Measured on Moscow
# iPhones in October 2026: 60-75% of "private" listings, raising the median
# by up to 10%. Private people repeat round prices (25 000) — those stay.
NETWORK_REPEATS = 4
NETWORK_WORDS = re.compile(r"запечатан[^.!\n]{0,80}гаранти|гаранти[^.!\n]{0,80}запечатан", re.IGNORECASE)
NETWORK = "сеть перекупов (одинаковые цены)"

# Locked or broken devices are sold far below market; they are marked and
# left out of the market price.
DEFECT_PATTERN = re.compile(
    r"(?<!не\s)бит(?:ый|ая|ое|ые)\b|заблок|блокировк|icloud|айклауд|аклауд|на\s+запчаст|на\s+разбор|под\s+восстановл|"
    r"не\s+включ|не\s+работает|разбит|треснут|утоплен|после\s+воды|залит|"
    r"не\s+ловит|плата\s+мертв|донор|неисправн|нерабоч|не\s+рабоч|списанн|на\s+детали",
    re.IGNORECASE,
)


@dataclass
class Offer:
    listing: apartments.Listing
    model: str
    memory: str
    seller: str
    delivery_only: bool
    avito_price_badge: str
    city: str = ""
    defect: bool = False
    network: bool = False
    market: int | None = None
    vs_market: float | None = None
    notes: list[str] = field(default_factory=list)

    @property
    def group(self) -> str:
        return f"{self.model}, {self.memory}" if self.memory else self.model


def offer_from_item(item: dict[str, Any]) -> Offer:
    listing = apartments.listing_from_item(item)
    parts = [part.strip() for part in listing.title.split(",") if part.strip()]
    model = parts[0] if parts else listing.title
    memory = next((part for part in parts[1:] if MEMORY_PATTERN.match(part)), "")
    badges = listing.avito_badges
    price_badge = next(
        (badge for badge in badges if "рыночн" in badge.lower()),
        "",
    )
    return Offer(
        listing=listing,
        model=model,
        memory=memory,
        seller=listing.seller_label or "частное лицо",
        delivery_only="Только доставка" in badges,
        avito_price_badge=price_badge,
        city=str(item.get("urlPath") or "").strip("/").split("/")[0],
        defect=bool(DEFECT_PATTERN.search(f"{listing.title} {listing.description}")),
    )


@dataclass(frozen=True)
class GroupStats:
    group: str
    count: int
    used: int
    low: int
    q1: int
    median: int
    q3: int
    high: int


def network_prices(offers: list[Offer]) -> Counter:
    """How many listings of each (group, price) private sellers without a shop history put up."""
    return Counter(
        (offer.group, offer.listing.price) for offer in offers
        if offer.listing.price and offer.seller == "частное лицо"
        and (offer.listing.closed_listings or 0) < SHOP_CLOSED_LISTINGS
    )


def is_network(offer: Offer, prices: Counter) -> bool:
    price = offer.listing.price
    if price and price % 1000 and prices[(offer.group, price)] >= NETWORK_REPEATS:
        return True
    return bool(NETWORK_WORDS.search(f"{offer.listing.title} {offer.listing.description}"))


def mark_networks(offers: list[Offer]) -> Counter:
    prices = network_prices(offers)
    for offer in offers:
        offer.network = is_network(offer, prices)
    return prices


def market_stats(offers: list[Offer]) -> dict[str, GroupStats]:
    """Market prices by group, leaving out broken devices and the reseller network (mark_networks)."""
    groups: dict[str, list[int]] = {}
    for offer in offers:
        if offer.listing.price and not offer.defect and not offer.network:
            groups.setdefault(offer.group, []).append(offer.listing.price)
    result = {}
    for group, prices in groups.items():
        if len(prices) < MIN_GROUP_SIZE:
            continue
        raw_median = statistics.median(prices)
        kept = sorted(
            price for price in prices
            if OUTLIER_LOW * raw_median <= price <= OUTLIER_HIGH * raw_median
        )
        if len(kept) < MIN_GROUP_SIZE:
            continue
        q1, median, q3 = statistics.quantiles(kept, n=4)
        result[group] = GroupStats(
            group=group, count=len(prices), used=len(kept),
            low=kept[0], q1=round(q1), median=round(median), q3=round(q3), high=kept[-1],
        )
    return result


def annotate(
    offers: list[Offer], stats: dict[str, GroupStats], home_city: str = ""
) -> None:
    for offer in offers:
        group = stats.get(offer.group)
        offer.notes = []
        if group and offer.listing.price:
            offer.market = group.median
            offer.vs_market = offer.listing.price / group.median - 1
            if offer.defect:
                offer.notes.append("заблокирован/неисправен (по описанию)")
            elif offer.listing.price < OUTLIER_LOW * group.median:
                offer.notes.append("подозрительно дёшево: битый, запчасти или обман?")
            elif offer.listing.price <= CHEAP_SHARE * group.median:
                offer.notes.append("дешевле рынка")
        if offer.seller != "частное лицо":
            offer.notes.append(f"продавец: {offer.seller}")
        elif (offer.listing.closed_listings or 0) >= SHOP_CLOSED_LISTINGS:
            offer.notes.append(
                f"похоже на магазин/перекупа: {offer.listing.closed_listings} завершённых"
            )
        if offer.delivery_only:
            offer.notes.append("только доставка")
        if home_city and offer.city and offer.city != home_city:
            offer.notes.append(f"другой город: {offer.city}")


HEADER_FILL = "1F4E78"
CHEAP_FILL = "C6EFCE"       # "К рынку": noticeably cheaper than the market
DEAR_FILL = "FFC7CE"        # "К рынку": noticeably dearer
DEFECT_FILL = "F8D7DA"      # the whole row: broken or locked by its description
SUSPICIOUS_FILL = "FFF2CC"  # the whole row: far below the market
PRICE_FORMAT = '#,##0" ₽"'
SHARE_FORMAT = "+0%;-0%;0%"
TOP_CHEAPEST = 10
# The cheapest list skips prices below this share of the private median:
# such "offers" are lots, smaller sizes, accessories or typos.
TOO_CHEAP_SHARE = 0.25
RESELLER = "частное лицо, похоже на перекупа"
# The seller column already says this; the table's notes do not repeat it.
SELLER_NOTES = ("продавец:", "похоже на магазин")

# (header, width, kind); a column with nothing in it is left out of the table.
LISTING_COLUMNS = (
    ("Новое", 8, "text"),
    ("Название", 46, "link"),
    ("Цена", 12, "price"),
    ("Рыночная цена", 13, "price"),
    ("К рынку", 9, "share"),
    ("Оценка Авито", 18, "text"),
    ("Продавец", 24, "text"),
    ("Завершённых объявлений у продавца", 14, "number"),
    ("Где", 30, "text"),
    ("Опубликовано", 14, "text"),
    ("Пометки", 40, "text"),
    ("С чем сравнивали цену", 30, "text"),
    ("Описание", 70, "text"),
)


def seller_kind(offer: Offer) -> str:
    if offer.network:
        return NETWORK
    if offer.seller != "частное лицо":
        return offer.seller
    if (offer.listing.closed_listings or 0) >= SHOP_CLOSED_LISTINGS:
        return RESELLER
    return "частное лицо"


def is_other_city(offer: Offer, home_city: str) -> bool:
    return bool(home_city and offer.city and offer.city != home_city)


def listing_row(offer: Offer) -> dict[str, Any]:
    listing = offer.listing
    return {
        "Новое": "новое" if listing.is_new else "",
        "Название": listing.title,
        "Цена": listing.price or None,  # 0 ₽ means "price on request"
        "Рыночная цена": offer.market,
        "К рынку": offer.vs_market,
        "Оценка Авито": offer.avito_price_badge,
        "Продавец": seller_kind(offer),
        "Завершённых объявлений у продавца": listing.closed_listings,
        "Где": apartments_short_address(listing),
        "Опубликовано": listing.published,
        "Пометки": "; ".join(note for note in offer.notes if not note.startswith(SELLER_NOTES)),
        "С чем сравнивали цену": offer.group if offer.market else "",
        "Описание": listing.description,
    }


def write_workbook(
    path: Path,
    offers: list[Offer],
    stats: dict[str, GroupStats],
    *,
    title: str = "",
    link: str = "",
    home_city: str = "",
) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    ordered = sorted(offers, key=lambda o: (not o.listing.price, o.listing.price or 0, o.listing.title))
    write_summary(workbook.active, ordered, stats, title=title, link=link, home_city=home_city)
    write_listings(workbook.create_sheet("Объявления"), ordered)
    write_market(workbook.create_sheet("Рыночные цены"), stats)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _header(cells: Any) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    for cell in cells:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _link(cell: Any, url: str) -> None:
    from openpyxl.styles import Font

    cell.hyperlink = url
    cell.font = Font(color="0563C1", underline="single")


def write_summary(
    sheet: Any, offers: list[Offer], stats: dict[str, GroupStats], *, title: str, link: str, home_city: str
) -> None:
    from datetime import datetime

    from openpyxl.styles import Alignment, Font

    sheet.title = "Сводка"
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 50
    sheet.column_dimensions["C"].width = 30
    sheet.column_dimensions["D"].width = 40
    row = 1

    def put(label: str, value: Any = None, *, number_format: str | None = None, bold: bool = False) -> None:
        nonlocal row
        sheet.cell(row=row, column=1, value=label).font = Font(bold=bold)
        if value is not None:
            cell = sheet.cell(row=row, column=2, value=value)
            cell.alignment = Alignment(horizontal="left")
            if number_format:
                cell.number_format = number_format
        row += 1

    sheet.cell(row=row, column=1, value=title or "Объявления Авито").font = Font(bold=True, size=14)
    row += 1
    if link:
        put("Поиск на Авито", "открыть")
        _link(sheet.cell(row=row - 1, column=2), link)
    put("Собрано", datetime.now().strftime("%d.%m.%Y %H:%M"))
    row += 1

    kinds = [seller_kind(offer) for offer in offers]
    new = sum(offer.listing.is_new for offer in offers)
    put("Всего объявлений", len(offers), bold=True)
    put("   от частных лиц", kinds.count("частное лицо"))
    put("   от частных лиц, похожих на перекупщиков", kinds.count(RESELLER))
    put("   сеть перекупов с одинаковыми ценами", kinds.count(NETWORK))
    put("   от магазинов и компаний",
        len(kinds) - kinds.count("частное лицо") - kinds.count(RESELLER) - kinds.count(NETWORK))
    if home_city:
        put("   из других городов (с доставкой)", sum(is_other_city(offer, home_city) for offer in offers))
    put("   неисправные или на запчасти (по описанию)", sum(offer.defect for offer in offers))
    if 0 < new < len(offers):
        put("   новых с прошлого сбора", new)
    row += 1

    priced = [offer for offer in offers if offer.listing.price and not offer.defect]
    private = sorted(offer.listing.price for offer in priced if seller_kind(offer) == "частное лицо")
    everyone = sorted(offer.listing.price for offer in priced)
    if everyone:
        put("Цены (без неисправных и без «цена не указана»)", bold=True)
        for column, name in enumerate(("", "Частные лица (без перекупщиков)", "Все продавцы"), start=1):
            sheet.cell(row=row, column=column, value=name)
        _header(sheet[row][:3])
        row += 1
        for label, pick in (
            ("самая низкая", lambda prices: prices[0]),
            ("медиана: половина дешевле, половина дороже", lambda prices: round(statistics.median(prices))),
            ("самая высокая", lambda prices: prices[-1]),
        ):
            sheet.cell(row=row, column=1, value=f"   {label}")
            for column, prices in ((2, private), (3, everyone)):
                if prices:
                    cell = sheet.cell(row=row, column=column, value=pick(prices))
                    cell.number_format = PRICE_FORMAT
                    cell.alignment = Alignment(horizontal="left")
            row += 1
        row += 1

    floor = TOO_CHEAP_SHARE * statistics.median(private) if private else 0
    cheapest = [
        offer for offer in priced
        if seller_kind(offer) == "частное лицо" and offer.listing.price >= floor
        and not is_other_city(offer, home_city)
        and not any(note.startswith("подозрительно дёшево") for note in offer.notes)
    ][:TOP_CHEAPEST]
    if cheapest:
        put(f"Самые дешёвые от частных лиц{' в своём городе' if home_city else ''}", bold=True)
        if floor:
            put(f"   без цен ниже {round(floor):,} ₽ (четверть медианы): это обычно лоты, мелкие объёмы или ошибки"
                .replace(",", " "))
        for column, name in enumerate(("Цена", "Название (ссылка)", "Где", "Пометки"), start=1):
            sheet.cell(row=row, column=column, value=name)
        _header(sheet[row][:4])
        row += 1
        for offer in cheapest:
            sheet.cell(row=row, column=1, value=offer.listing.price).number_format = PRICE_FORMAT
            _link(sheet.cell(row=row, column=2, value=offer.listing.title), offer.listing.url)
            sheet.cell(row=row, column=3, value=apartments_short_address(offer.listing))
            sheet.cell(row=row, column=4, value=listing_row(offer)["Пометки"])
            row += 1
        row += 1

    put("Как читать таблицу", bold=True)
    for line in (
        "Лист «Объявления» — все объявления, сначала самые дешёвые.",
        "Стрелки в заголовках — фильтры: например, оставить только «частное лицо».",
        "Название — ссылка: клик открывает объявление на Авито.",
        "«К рынку»: −20% значит на 20% дешевле рыночной цены (зелёный), +20% — дороже (красный).",
        "Рыночная цена — медиана похожих объявлений (та же модель), см. лист «Рыночные цены».",
        "Красная строка — по описанию неисправен или на запчасти; жёлтая — подозрительно дёшево.",
        f"«Похоже на перекупа» — частное лицо с {SHOP_CLOSED_LISTINGS}+ завершёнными объявлениями.",
        "«Сеть перекупов» — аккаунты без истории с одной и той же некруглой ценой (34 990) "
        "или «запечатанный, с гарантией»; в рыночную цену не входят.",
    ):
        put("   " + line)
    if not stats:
        put("   Похожих объявлений (≥3 одной модели) не нашлось, поэтому рыночной цены нет.")


def write_listings(sheet: Any, offers: list[Offer]) -> None:
    from openpyxl.styles import PatternFill
    from openpyxl.utils import get_column_letter

    rows = [listing_row(offer) for offer in offers]
    columns = [
        column for column in LISTING_COLUMNS
        if any(row[column[0]] not in (None, "") for row in rows)
        # On a first collection every listing is new: the column says nothing.
        and not (column[0] == "Новое" and all(offer.listing.is_new for offer in offers))
    ]
    sheet.append([name for name, _, _ in columns])
    _header(sheet[1])
    sheet.row_dimensions[1].height = 32
    for index, (_, width, _) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for offer, values in zip(offers, rows):
        sheet.append([values[name] for name, _, _ in columns])
        cells = sheet[sheet.max_row]
        fill = None
        if offer.defect:
            fill = DEFECT_FILL
        elif any(note.startswith("подозрительно дёшево") for note in offer.notes):
            fill = SUSPICIOUS_FILL
        for cell, (name, _, kind) in zip(cells, columns):
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
            if kind == "price":
                cell.number_format = PRICE_FORMAT
            elif kind == "share" and cell.value is not None:
                cell.number_format = SHARE_FORMAT
                if not fill and cell.value <= CHEAP_SHARE - 1:
                    cell.fill = PatternFill("solid", fgColor=CHEAP_FILL)
                elif not fill and cell.value >= 1 - CHEAP_SHARE:
                    cell.fill = PatternFill("solid", fgColor=DEAR_FILL)
            elif kind == "link":
                _link(cell, offer.listing.url)
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions


def write_market(sheet: Any, stats: dict[str, GroupStats]) -> None:
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    sheet.cell(row=1, column=1, value=(
        "Рыночная цена — медиана цен похожих объявлений (одна модель). Неисправные и слишком "
        "далёкие от остальных цены (битые, запчасти, опечатки) в расчёт не входят."
    )).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    sheet.row_dimensions[1].height = 32
    headers = ["Модель", "Объявлений", "Учтено в расчёте", "Самая низкая", "Дешёвая четверть — до",
               "Рыночная цена (медиана)", "Дорогая четверть — от", "Самая высокая"]
    for column, name in enumerate(headers, start=1):
        sheet.cell(row=3, column=column, value=name)
    _header(sheet[3])
    sheet.row_dimensions[3].height = 32
    for index, width in enumerate((40, 12, 12, 13, 14, 14, 14, 13), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for group in sorted(stats.values(), key=lambda s: (-s.count, s.group)):
        sheet.append([group.group, group.count, group.used, group.low, group.q1, group.median, group.q3, group.high])
        for cell in sheet[sheet.max_row][3:]:
            cell.number_format = PRICE_FORMAT
        sheet.cell(row=sheet.max_row, column=6).font = Font(bold=True)
    if not stats:
        sheet.cell(row=4, column=1, value="Похожих объявлений (≥3 одной модели) не нашлось.")
    sheet.freeze_panes = "B4"


def apartments_short_address(listing: apartments.Listing) -> str:
    return re.sub(r"^[^,]+,\s*", "", listing.address) if "," in listing.address else listing.address


def price_text(value: int | None) -> str:
    return f"{value:,} ₽".replace(",", " ") if value else "—"


def alert_fields(offer: Offer) -> dict[str, Any]:
    listing = offer.listing
    return {
        "id": listing.id,
        "title": listing.title,
        "price": price_text(listing.price),
        "market": price_text(offer.market),
        "vs_market": f"{offer.vs_market:+.0%}" if offer.vs_market is not None else "",
        "avito_badge": offer.avito_price_badge,
        "seller": offer.seller,
        "closed_listings": listing.closed_listings,
        "notes": offer.notes,
        "published": listing.published,
        "address": apartments_short_address(listing),
        "url": listing.url,
    }
